"""Export router for XLSX and marked PDF generation."""
import json
import os
from pathlib import Path
from typing import Optional

from fastapi import APIRouter, HTTPException, BackgroundTasks
from fastapi.responses import FileResponse

router = APIRouter(prefix="/export", tags=["export"])

def _get_base_dir():
    from ..config_manager import get_config_value
    return Path(get_config_value("paths.base_dir", "."))


def _get_project_paths(project_id: str):
    """Get prediction and evaluation file paths for a project."""
    base = _get_base_dir()
    pred_path = base / "outputs" / project_id / "prediction.json"
    eval_path = base / "outputs" / project_id / "evaluation_report.json"
    return pred_path, eval_path


@router.get("/xlsx/{project_id}")
def export_xlsx(project_id: str, background_tasks: BackgroundTasks):
    """Export prediction to Excel (.xlsx) file."""
    pred_path, eval_path = _get_project_paths(project_id)
    
    if not pred_path.exists():
        raise HTTPException(status_code=404, detail="No prediction found for this project")
    
    try:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    except ImportError:
        raise HTTPException(status_code=500, detail="openpyxl not installed")
    
    with open(pred_path, 'r', encoding='utf-8') as f:
        data = json.load(f)
    
    evaluation = None
    if eval_path.exists():
        with open(eval_path, 'r', encoding='utf-8') as f:
            evaluation = json.load(f)
    
    wb = openpyxl.Workbook()
    
    # Summary sheet
    ws_summary = wb.active
    ws_summary.title = "Summary"
    
    # Header styling
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
    thin_border = Border(
        left=Side(style='thin'), right=Side(style='thin'),
        top=Side(style='thin'), bottom=Side(style='thin')
    )
    
    ws_summary.append(["AI Takeoff Builder - Flooring Challenge Export"])
    ws_summary.append(["Project ID:", project_id])
    ws_summary.append([])
    
    if evaluation:
        matched = evaluation.get('matched_items', 0)
        missing = len(evaluation.get('missing_items', []))
        extra = len(evaluation.get('extra_items', []))
        total_expected = matched + missing
        accuracy = (matched / total_expected * 100) if total_expected > 0 else 0
        
        ws_summary.append(["Evaluation Results"])
        ws_summary.append(["Matched Items:", matched])
        ws_summary.append(["Missing Items:", missing])
        ws_summary.append(["Extra Items:", extra])
        ws_summary.append(["Accuracy:", f"{accuracy:.1f}%"])
        ws_summary.append(["Target:", "75%"])
        ws_summary.append([])
    
    ws_summary.append(["Total Line Items:", data.get('total_line_items', 0)])
    ws_summary.append(["Trades:", ", ".join(data.get('trades', []))])
    
    # Auto-size columns
    for col in ws_summary.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = min(max_length + 2, 60)
        ws_summary.column_dimensions[column].width = adjusted_width
    
    # Line Items sheet
    ws_items = wb.create_sheet("Line Items")
    headers = ["#", "Description", "Trade", "Quantity", "Unit", "Confidence", "Source Reference"]
    ws_items.append(headers)
    
    for cell in ws_items[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal='center')
        cell.border = thin_border
    
    line_items = data.get('line_items', [])
    for idx, item in enumerate(line_items, 1):
        ws_items.append([
            idx,
            item.get('description', ''),
            item.get('trade', ''),
            item.get('quantity', ''),
            item.get('unit', ''),
            item.get('confidence', ''),
            item.get('source_reference', ''),
        ])
    
    # Style data rows
    for row in ws_items.iter_rows(min_row=2, max_row=ws_items.max_row):
        for cell in row:
            cell.border = thin_border
            cell.alignment = Alignment(vertical='top', wrap_text=True)
    
    # Auto-size columns for items
    for col in ws_items.columns:
        max_length = 0
        column = col[0].column_letter
        for cell in col:
            try:
                if len(str(cell.value)) > max_length:
                    max_length = len(str(cell.value))
            except:
                pass
        adjusted_width = min(max_length + 2, 60)
        ws_items.column_dimensions[column].width = adjusted_width
    
    # Missing items sheet (if evaluation exists)
    if evaluation and evaluation.get('missing_items'):
        ws_missing = wb.create_sheet("Missing Items")
        ws_missing.append(["#", "Expected Description"])
        for cell in ws_missing[1]:
            cell.font = header_font
            cell.fill = header_fill
        for idx, item in enumerate(evaluation['missing_items'], 1):
            ws_missing.append([idx, item if isinstance(item, str) else item.get('description', '')])
    
    # Extra items sheet (if evaluation exists)
    if evaluation and evaluation.get('extra_items'):
        ws_extra = wb.create_sheet("Extra Items")
        ws_extra.append(["#", "Predicted Description"])
        for cell in ws_extra[1]:
            cell.font = header_font
            cell.fill = PatternFill(start_color="DC2626", end_color="DC2626", fill_type="solid")
        for idx, item in enumerate(evaluation['extra_items'], 1):
            ws_extra.append([idx, item if isinstance(item, str) else item.get('description', '')])
    
    output_path = _get_base_dir() / "outputs" / project_id / f"{project_id}_takeoff.xlsx"
    wb.save(output_path)
    
    return FileResponse(
        path=output_path,
        filename=f"{project_id}_takeoff.xlsx",
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )


@router.get("/marked-pdf/{project_id}")
def export_marked_pdf(project_id: str, background_tasks: BackgroundTasks):
    """Generate a marked PDF with extracted item annotations."""
    pred_path, _ = _get_project_paths(project_id)
    
    if not pred_path.exists():
        raise HTTPException(status_code=404, detail="No prediction found for this project")
    
    try:
        import fitz
    except ImportError:
        raise HTTPException(status_code=500, detail="PyMuPDF not installed")
    
    with open(pred_path, 'r', encoding='utf-8') as f:
        data = json.load(f)
    
    line_items = data.get('line_items', [])
    
    # Find the original PDF file(s)
    base = _get_base_dir()
    pdf_files = []
    for search_dir in [base / "assessment2", base / "client_files", base / "tmp"]:
        if search_dir.exists():
            for p in search_dir.rglob("*.pdf"):
                if project_id in str(p):
                    pdf_files.append(p)
    
    if not pdf_files:
        raise HTTPException(status_code=404, detail="No source PDF files found for this project")
    
    # Use the largest PDF (likely the main drawing)
    source_pdf = max(pdf_files, key=lambda p: p.stat().st_size)
    
    doc = fitz.open(str(source_pdf))
    
    # Add a cover page with summary
    cover = doc.new_page(-1, width=612, height=792)
    y = 60
    
    cover.insert_text((60, y), "AI Takeoff Builder - Marked Drawing", fontsize=20, fontname="helv", color=(0, 0, 0.5))
    y += 40
    cover.insert_text((60, y), f"Project: {project_id}", fontsize=14, fontname="helv")
    y += 30
    cover.insert_text((60, y), f"Total Items Extracted: {len(line_items)}", fontsize=12, fontname="helv")
    y += 25
    cover.insert_text((60, y), f"Source: {source_pdf.name}", fontsize=10, fontname="helv", color=(0.4, 0.4, 0.4))
    y += 40
    
    # List items on cover
    cover.insert_text((60, y), "Extracted Line Items:", fontsize=12, fontname="helv", color=(0, 0, 0.5))
    y += 20
    
    for idx, item in enumerate(line_items[:50], 1):
        desc = item.get('description', '')[:80]
        trade = item.get('trade', '')
        text = f"{idx}. [{trade}] {desc}"
        cover.insert_text((60, y), text, fontsize=8, fontname="helv")
        y += 12
        if y > 750:
            break
    
    # Search and highlight flooring-related keywords on each page
    flooring_keywords = []
    for item in line_items:
        desc = item.get('description', '').upper()
        # Extract key terms from description
        terms = []
        for term in ['PORCELAIN', 'VINYL', 'VCT', 'CARPET', 'HARDWOOD', 'SCHLUTER', 'RUBBER BASE', 'MDF BASE', 'WOOD BASE', 'WOOD LEDGER']:
            if term in desc:
                terms.append(term)
        if terms:
            flooring_keywords.append({
                'terms': terms,
                'description': item.get('description', '')[:60],
                'trade': item.get('trade', ''),
            })
    
    # Highlight pages
    for page_num in range(len(doc) - 1):  # Skip cover
        page = doc[page_num]
        
        for kw in flooring_keywords:
            for term in kw['terms']:
                rects = page.search_for(term)
                for rect in rects:
                    # Add highlight annotation
                    highlight = page.add_highlight_annot(rect)
                    highlight.set_colors(stroke=(1, 1, 0))  # Yellow
                    highlight.update()
                    
                    # Add a small text annotation near the highlight
                    annot_rect = fitz.Rect(rect.x0, rect.y0 - 12, rect.x1, rect.y0)
                    annot = page.add_text_annot(annot_rect.tl, kw['description'])
                    annot.set_colors(stroke=(0, 0, 0.8))
                    annot.update()
    
    output_path = _get_base_dir() / "outputs" / project_id / f"{project_id}_marked.pdf"
    doc.save(output_path)
    doc.close()
    
    return FileResponse(
        path=output_path,
        filename=f"{project_id}_marked.pdf",
        media_type="application/pdf"
    )
